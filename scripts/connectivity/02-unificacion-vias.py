import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ARCHIVO = r"C:\A_GS1_PROYECTOS\0_Documents_gs\database\output\02-raw_vias.xlsx"

COLUMNAS_OBLIGATORIAS = [
    "HOJA_ORIGEN", "SALON", "FILA", "BASTIDOR", "EQUIPO",
    "INTERFAZ-PUERTO-GE", "ANILLO", "PACHEO", "TRONCAL-CABLE",
    "ODF CALLE", "HILO", "OBSERVACION"
]

COLUMNAS_ESTANDAR = [c for c in COLUMNAS_OBLIGATORIAS if c != "HOJA_ORIGEN"]

HOJAS_EXCLUIDAS = {"metro_san carlos", "UNIFICADO"}

MAPEO_COLUMNAS = {
    "INTERFAZ-PUERTO-": "INTERFAZ-PUERTO-GE",
    "INTERFAZ-PUERTO": "INTERFAZ-PUERTO-GE",
    "INTERFAZ/PUERTO/GE": "INTERFAZ-PUERTO-GE",
    "ODF PACHEO": "PACHEO",
    "ODF PATCHEO": "PACHEO",
    "PATCHEO": "PACHEO",
    "OBSERVACIONES": "OBSERVACION",
}

COLUMNAS_A_ELIMINAR = {"VIAS"}


def normalizar_columnas(df):
    df.columns = [
        MAPEO_COLUMNAS.get(str(c).strip().upper(), str(c).strip().upper())
        for c in df.columns
    ]
    return df


def forzar_headers_string(df):
    nuevas, vistos = [], {}
    for c in df.columns:
        nombre = str(c).strip()
        if nombre == "" or nombre.lower() == "nan":
            nombre = "COLUMNA"
        if nombre in vistos:
            vistos[nombre] += 1
            nombre = f"{nombre}_{vistos[nombre]}"
        else:
            vistos[nombre] = 0
        nuevas.append(nombre)
    df.columns = nuevas
    return df


def eliminar_columnas_inutiles(df):
    cols_unnamed = [c for c in df.columns if str(c).upper().startswith("UNNAMED")]
    df = df.drop(columns=cols_unnamed, errors="ignore")

    cols_basura = [c for c in df.columns if str(c).strip().upper() in COLUMNAS_A_ELIMINAR]
    df = df.drop(columns=cols_basura, errors="ignore")

    return df


def corregir_cruces(df):
    orden_real = [
        "SALON", "FILA", "BASTIDOR", "EQUIPO",
        "INTERFAZ-PUERTO-GE", "ANILLO", "PACHEO",
        "TRONCAL-CABLE", "ODF CALLE", "HILO", "OBSERVACION",
    ]
    df = df.iloc[:, :len(orden_real)].copy()
    df.columns = orden_real
    return df


def corregir_san_jose(df):
    cols = df.columns.tolist()
    idx_pacheo = next(
        (i for i, c in enumerate(cols) if str(c).strip().upper() == "ODF PACHEO"),
        None,
    )

    if idx_pacheo is not None and idx_pacheo + 1 < len(cols):
        col_pacheo = df.columns[idx_pacheo]
        col_hilo_espurio = df.columns[idx_pacheo + 1]

        pacheo_vals = df[col_pacheo].astype(str).str.strip()
        hilo_vals = df[col_hilo_espurio].astype(str).str.strip()

        vacio = (hilo_vals == "") | (hilo_vals == "nan")
        nueva_col = pacheo_vals + "-" + hilo_vals
        nueva_col[vacio] = pacheo_vals[vacio]

        df[col_pacheo] = nueva_col
        df = df.drop(columns=[col_hilo_espurio])

    df = df.rename(columns={"HILO.1": "HILO"})
    return normalizar_columnas(df)


def colapsar_extras(df):
    """
    Une todas las columnas que NO son obligatorias en una sola columna
    llamada EXTRAS, separando valores con '°'.
    Debe llamarse ANTES de seleccionar solo las columnas obligatorias.
    """
    no_extras = set(COLUMNAS_OBLIGATORIAS)
    extras = [c for c in df.columns if c not in no_extras]

    if not extras:
        df["EXTRAS"] = ""
        return df

    def unir_fila(row):
        valores = []
        for col in extras:
            v = row[col]
            if v is None:
                continue
            if isinstance(v, float) and pd.isna(v):
                continue
            v_str = str(v).strip()
            if v_str == "" or v_str.lower() in {"nan", "none", "nat"}:
                continue
            valores.append(v_str)
        return "°".join(valores)

    df["EXTRAS"] = df[extras].apply(unir_fila, axis=1)
    df = df.drop(columns=extras)
    return df


def procesar_hoja(nombre_hoja, df_raw):
    df = df_raw.dropna(axis=1, how="all").dropna(axis=0, how="all")

    if nombre_hoja == "SAN JOSE":
        df = corregir_san_jose(df)
    else:
        df = normalizar_columnas(df)

    if nombre_hoja == "CRUCES":
        df = corregir_cruces(df)

    df = eliminar_columnas_inutiles(df)
    df = forzar_headers_string(df)

    # 1) Rellenar columnas obligatorias faltantes (antes de colapsar)
    for col in COLUMNAS_ESTANDAR:
        if col not in df.columns:
            df[col] = np.nan

    # 2) Colapsar extras (mientras aún existen)
    df = colapsar_extras(df)

    # 3) Insertar HOJA_ORIGEN si no está
    if "HOJA_ORIGEN" not in df.columns:
        df.insert(0, "HOJA_ORIGEN", nombre_hoja)
    else:
        df["HOJA_ORIGEN"] = nombre_hoja

    # 4) Reordenar columnas finales
    orden_final = COLUMNAS_OBLIGATORIAS + ["EXTRAS"]
    df = df[orden_final]

    return df


def leer_todo(archivo):
    try:
        return pd.read_excel(archivo, sheet_name=None, engine="calamine")
    except Exception:
        return pd.read_excel(archivo, sheet_name=None, engine="openpyxl")


def unificar(archivo):
    hojas = leer_todo(archivo)

    dataframes = [
        procesar_hoja(nombre, df)
        for nombre, df in hojas.items()
        if nombre not in HOJAS_EXCLUIDAS
    ]

    return pd.concat(dataframes, ignore_index=True, sort=False)


def escribir_todo_xlsxwriter(archivo, hojas_originales, df_final,
                             nombre_hoja="UNIFICADO"):
    with pd.ExcelWriter(archivo, engine="xlsxwriter") as writer:
        df_final.to_excel(writer, index=False, sheet_name=nombre_hoja)
        for nombre, df in hojas_originales.items():
            if nombre == "metro_san carlos":
                continue
            df.to_excel(writer, index=False, sheet_name=nombre)


def aplicar_tabla_openpyxl(archivo, nombre_hoja="UNIFICADO",
                           nombre_tabla="TablaUnificada"):
    wb = load_workbook(archivo)
    ws = wb[nombre_hoja]
    ws.tables.clear()

    max_col = ws.max_column
    max_row = ws.max_row
    col_fin = get_column_letter(max_col)
    rango = f"A1:{col_fin}{max_row}"

    tabla = Table(displayName=nombre_tabla, ref=rango)
    tabla.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tabla)

    wb.save(archivo)
    wb.close()


if __name__ == "__main__":
    hojas_originales = leer_todo(ARCHIVO)
    df_final = unificar(ARCHIVO)

    print(f"Filas unificadas: {len(df_final):,}")
    print(f"Hojas procesadas: {df_final['HOJA_ORIGEN'].nunique()}")
    print(f"Columnas finales: {list(df_final.columns)}")
    print(f"Filas con EXTRAS no vacías: {(df_final['EXTRAS'] != '').sum():,}")
    print(df_final["HOJA_ORIGEN"].value_counts())

    escribir_todo_xlsxwriter(ARCHIVO, hojas_originales, df_final)
    aplicar_tabla_openpyxl(ARCHIVO)

    print(f"Archivo reescrito con 'UNIFICADO' al inicio y Tabla aplicada: {ARCHIVO}")